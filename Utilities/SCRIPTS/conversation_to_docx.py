"""
Convert conversation chunks to clean DOCX format.
Handles broken JSON, strips brackets, creates Word documents.

Usage:
    python conversation_to_docx.py <chunks_directory>
"""

import re
import sys
import argparse
from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

# Fix Windows console encoding
if sys.platform == 'win32':
    import codecs
    sys.stdout = codecs.getwriter('utf-8')(sys.stdout.buffer, 'ignore')
    sys.stderr = codecs.getwriter('utf-8')(sys.stderr.buffer, 'ignore')


class ConversationToDOCX:
    """Convert conversations to DOCX, handling broken JSON."""

    def __init__(self, chunks_dir: str, output_prefix: str = "perplexity_chats"):
        self.chunks_dir = Path(chunks_dir)
        self.output_prefix = output_prefix
        self.output_dir = self.chunks_dir.parent / f"{output_prefix}_docx"
        self.output_dir.mkdir(exist_ok=True)

    def _strip_brackets(self, text: str) -> str:
        """Remove all brackets from text."""
        # Remove brackets but keep content
        text = re.sub(r'[\{\}\[\]\(\)]', '', text)
        return text

    def _clean_text(self, text: str) -> str:
        """Clean and unescape text."""
        # Unescape JSON
        text = text.replace('\\n', '\n')
        text = text.replace('\\"', '"')
        text = text.replace('\\/', '/')
        text = text.replace('\\\\', '\\')
        text = text.replace('\\t', '\t')

        # Strip brackets
        text = self._strip_brackets(text)

        # Clean whitespace
        lines = [line.rstrip() for line in text.split('\n')]
        text = '\n'.join(lines)

        # Remove excessive blank lines
        while '\n\n\n' in text:
            text = text.replace('\n\n\n', '\n\n')

        return text.strip()

    def _extract_between(self, text: str, start_marker: str, end_marker: str = None) -> str:
        """Extract text between markers."""
        try:
            start_idx = text.find(start_marker)
            if start_idx == -1:
                return ""

            start_idx += len(start_marker)

            if end_marker:
                end_idx = text.find(end_marker, start_idx)
                if end_idx == -1:
                    return text[start_idx:].strip()
                return text[start_idx:end_idx].strip()
            else:
                return text[start_idx:].strip()
        except:
            return ""

    def _find_all_patterns(self, text: str, pattern: str) -> list:
        """Find all matches of a pattern."""
        matches = []
        for match in re.finditer(pattern, text, re.DOTALL):
            matches.append(match.group(1) if match.groups() else match.group(0))
        return matches

    def _extract_conversations(self, content: str) -> list:
        """Extract conversations using fuzzy pattern matching."""
        conversations = []

        # Split by context_uuid as marker (each conversation starts with one)
        parts = re.split(r'"context_uuid"\s*:', content)

        for part in parts[1:]:  # Skip first empty part
            try:
                # Extract title
                title_match = re.search(r'"context_title"\s*:\s*"([^"]*(?:\\.[^"]*)*)"', part)
                title = self._clean_text(title_match.group(1)) if title_match else "Untitled"

                # Extract created date
                created_match = re.search(r'"created_at"\s*:\s*"([^"]+)"', part)
                created = created_match.group(1) if created_match else ""

                # Extract entries - look for query/answer OR role/content patterns
                entries = []

                # Try query/answer pattern first (Perplexity format)
                qa_matches = re.finditer(
                    r'"query"\s*:\s*"([^"]*(?:\\.[^"]*)*)".*?"answer"\s*:\s*"([^"]*(?:\\.[^"]*)*)"',
                    part,
                    re.DOTALL
                )

                for match in qa_matches:
                    query = self._clean_text(match.group(1))
                    answer = self._clean_text(match.group(2))

                    if query:
                        entries.append({'role': 'user', 'content': query})
                    if answer:
                        entries.append({'role': 'assistant', 'content': answer})

                # If no query/answer, try role/content pattern
                if not entries:
                    role_matches = re.finditer(
                        r'"role"\s*:\s*"(user|assistant)".*?"content"\s*:\s*"([^"]*(?:\\.[^"]*)*)"',
                        part,
                        re.DOTALL
                    )

                    for match in role_matches:
                        role = match.group(1)
                        content_text = self._clean_text(match.group(2))

                        if content_text and len(content_text) > 2:
                            entries.append({
                                'role': role,
                                'content': content_text
                            })

                if entries:
                    conversations.append({
                        'title': title,
                        'created': created,
                        'entries': entries
                    })

            except Exception as e:
                continue

        return conversations

    def _create_docx(self, conversations: list, output_file: Path, chunk_num: int):
        """Create a DOCX file from conversations."""
        doc = Document()

        # Set document margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)

        # Title
        title = doc.add_heading(f'{self.output_prefix.replace("_", " ").title()} - Part {chunk_num}', 0)
        title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # Summary
        p = doc.add_paragraph()
        p.add_run(f'Contains {len(conversations)} conversations').bold = True
        p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        doc.add_paragraph('_' * 80)

        # Add conversations
        for i, conv in enumerate(conversations, 1):
            # Conversation title
            heading = doc.add_heading(f"Conversation {i}: {conv['title']}", level=1)

            # Metadata
            if conv['created']:
                meta = doc.add_paragraph()
                meta.add_run(f"Date: {conv['created'][:10]}").italic = True

            doc.add_paragraph('─' * 60)

            # Add entries
            for entry in conv['entries']:
                role = entry['role']
                content = entry['content']

                # Role header
                p = doc.add_paragraph()
                run = p.add_run(f"{'User' if role == 'user' else 'Assistant'}:")
                run.bold = True
                run.font.size = Pt(11)

                if role == 'user':
                    run.font.color.rgb = RGBColor(0, 102, 204)  # Blue
                else:
                    run.font.color.rgb = RGBColor(34, 139, 34)  # Green

                # Content
                content_p = doc.add_paragraph(content)
                content_p.paragraph_format.left_indent = Inches(0.5)

                doc.add_paragraph()  # Spacing

            # Separator between conversations
            doc.add_paragraph('═' * 60)
            doc.add_paragraph()

        # Save
        doc.save(str(output_file))

    def process_chunk(self, chunk_file: Path, chunk_num: int):
        """Process a single chunk file."""
        print(f"📄 Processing {chunk_file.name}...", end=' ')

        try:
            # Read as raw text
            with open(chunk_file, 'r', encoding='utf-8', errors='ignore') as f:
                content = f.read()

            # Extract conversations
            conversations = self._extract_conversations(content)

            if not conversations:
                print("⚠️  No conversations found")
                return 0

            # Create DOCX
            output_file = self.output_dir / f"{self.output_prefix}_{chunk_num:04d}.docx"
            self._create_docx(conversations, output_file, chunk_num)

            file_size = output_file.stat().st_size / 1024
            print(f"✓ {len(conversations)} conversations ({file_size:.1f} KB)")

            return len(conversations)

        except Exception as e:
            print(f"❌ Error: {e}")
            return 0

    def process_all(self):
        """Process all chunk files."""
        # Find chunk files
        chunk_files = list(self.chunks_dir.glob('chunk_*.json'))
        if not chunk_files:
            chunk_files = list(self.chunks_dir.glob('*.json'))

        if not chunk_files:
            print(f"❌ No files found in {self.chunks_dir}")
            return

        chunk_files = sorted(chunk_files)

        print(f"\n📝 Converting to DOCX (Word format)")
        print(f"🔧 Robust mode: handles broken JSON, strips brackets")
        print(f"📁 Input: {self.chunks_dir}")
        print(f"📁 Output: {self.output_dir}")
        print(f"📄 Files: {len(chunk_files)}")
        print(f"\n{'='*60}")

        total_conversations = 0

        for i, chunk_file in enumerate(chunk_files, 1):
            count = self.process_chunk(chunk_file, i)
            total_conversations += count

        print(f"{'='*60}")
        print(f"✅ Conversion complete!")
        print(f"📊 Total conversations: {total_conversations}")
        print(f"📂 DOCX files saved to: {self.output_dir}")
        print(f"{'='*60}\n")


def main():
    parser = argparse.ArgumentParser(
        description='Convert conversation chunks to clean DOCX files',
        epilog="""
Examples:
  python conversation_to_docx.py ./chunks
  python conversation_to_docx.py ./simple_chunks --output-prefix "perplexity_chats"
        """
    )

    parser.add_argument('chunks_dir', help='Directory with chunk files')
    parser.add_argument(
        '--output-prefix',
        type=str,
        default='perplexity_chats',
        help='Prefix for output files (default: perplexity_chats)'
    )

    args = parser.parse_args()

    try:
        converter = ConversationToDOCX(
            chunks_dir=args.chunks_dir,
            output_prefix=args.output_prefix
        )
        converter.process_all()
    except Exception as e:
        print(f"\n❌ Error: {e}", file=sys.stderr)
        import traceback
        traceback.print_exc()
        sys.exit(1)


if __name__ == '__main__':
    main()
